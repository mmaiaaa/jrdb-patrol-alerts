#!/usr/bin/env python3
"""
Final experiment-closure audit for the JRDB supported-depth-band study.

This script DOES NOT rerun inference, alter thresholds, open new data,
or modify experimental outputs. It verifies the frozen provenance and
internal consistency of the completed study.
"""

import ast
import csv
import hashlib
import json
import subprocess
import sys
from collections import Counter
from pathlib import Path

ROOT = Path.cwd()

FREEZE_TAG = "jrdb-method-freeze-2026-09-12"
FREEZE_COMMIT = "3f9a8daadb832a6c6ea9251869b83be2a28c2746"

EXPECTED = {
    "reviewer_xlsx_sha256":
        "ceff3fbb2187c96d0fc3582792d0e6eaaadece2a4e71179c1f7916846a97e9b8",

    "yolo_weight_sha256":
        "f59b3d833e2ff32e194b5bb8e08d211dc7c5bdf144b90d2c8412c47ccfc83b36",

    "cubberly_prediction_sha256":
        "2e13fa212e647561765f0cc203ec60fd658e6377c6a7c655200700beddc56a16",

    "memorial_prediction_sha256":
        "d2d86947287c3e05654c8ec9913445990bebcb3b36b1b9cc841e1d593df36f13",

    "memorial_candidate_sha256":
        "8d3bd65efdc9fc1033cc611a38c67653ada8971a0eef1b9f5332847aabce9e51",

    "memorial_reporting_sha256":
        "d23e899bf7a7fa9e36b51057de1e381a0b1fba81a1092da08a760696d1bfdd10",

    "memorial_label_3d_sha256":
        "abb7befbb8949d307d12bcf38187787cc4928f2a1572327f2b79c71c71d1a723",

    "memorial_2d_combined_sha256":
        "5e065fcd1b1a2d92348e3b94c5f09fe0b2e6e29d6ed3e005bcf29e1a616042ab",
}


passes = []
warnings = []
failures = []


def sha256(path):
    h = hashlib.sha256()

    with path.open("rb") as f:
        for chunk in iter(
            lambda: f.read(1024 * 1024),
            b"",
        ):
            h.update(chunk)

    return h.hexdigest()


def git(*args, check=True):
    p = subprocess.run(
        ["git", *args],
        text=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
    )

    if check and p.returncode != 0:
        raise RuntimeError(
            f"git {' '.join(args)} failed:\n{p.stderr}"
        )

    return p


def PASS(name, detail=""):
    passes.append(name)
    suffix = f" — {detail}" if detail else ""
    print(f"[PASS] {name}{suffix}")


def WARN(name, detail=""):
    warnings.append(name)
    suffix = f" — {detail}" if detail else ""
    print(f"[WARN] {name}{suffix}")


def FAIL(name, detail=""):
    failures.append(name)
    suffix = f" — {detail}" if detail else ""
    print(f"[FAIL] {name}{suffix}")


def check(name, condition, detail=""):
    if condition:
        PASS(name, detail)
    else:
        FAIL(name, detail)


def require_file(path):
    if not path.is_file():
        FAIL("Required file exists", str(path))
        return False

    PASS("Required file exists", str(path))
    return True


def read_json(path):
    return json.loads(path.read_text(encoding="utf-8"))


def function_ast(path, function_name):
    tree = ast.parse(
        path.read_text(encoding="utf-8")
    )

    for node in tree.body:
        if (
            isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef))
            and node.name == function_name
        ):
            return ast.dump(
                node,
                annotate_fields=True,
                include_attributes=False,
            )

    raise ValueError(
        f"{function_name} not found in {path}"
    )


def literal_constant(path, name):
    tree = ast.parse(
        path.read_text(encoding="utf-8")
    )

    for node in tree.body:
        if not isinstance(node, ast.Assign):
            continue

        for target in node.targets:
            if (
                isinstance(target, ast.Name)
                and target.id == name
            ):
                try:
                    return ast.literal_eval(node.value)
                except Exception:
                    return None

    return None


print("=" * 78)
print("FINAL JRDB RESEARCH SANITY CHECK")
print("=" * 78)

# =====================================================================
# 1. Git provenance and method freeze
# =====================================================================

print("\n--- 1. GIT / METHOD FREEZE ---")

p = git(
    "rev-parse",
    f"{FREEZE_TAG}^{{commit}}",
    check=False,
)

if p.returncode == 0:
    actual = p.stdout.strip()

    check(
        "Method-freeze tag resolves to expected commit",
        actual == FREEZE_COMMIT,
        actual,
    )
else:
    FAIL(
        "Method-freeze tag exists",
        FREEZE_TAG,
    )


ancestor = git(
    "merge-base",
    "--is-ancestor",
    FREEZE_COMMIT,
    "HEAD",
    check=False,
)

check(
    "Current work descends from method freeze",
    ancestor.returncode == 0,
)


critical_frozen_files = [
    "scripts/cache_jrdb_native_detections.py",
    "scripts/cache_jrdb_dev_sensor_depth_candidates.py",
    "scripts/evaluate_jrdb_dev_surface_candidates_offline.py",
    "scripts/audit_jrdb_cubberly_projection.py",
    "scripts/audit_jrdb_dev_box_geometry.py",
    "scripts/audit_jrdb_box_depth_support.py",
    "scripts/review_jrdb_native_detections.py",
]

for filename in critical_frozen_files:
    if not Path(filename).exists():
        FAIL(
            "Frozen method source exists",
            filename,
        )
        continue

    diff = git(
        "diff",
        "--quiet",
        FREEZE_COMMIT,
        "--",
        filename,
        check=False,
    )

    check(
        "Frozen method source unchanged",
        diff.returncode == 0,
        filename,
    )


tracked_raw = git(
    "ls-files",
    "data",
    "outputs",
).stdout.strip()

check(
    "Raw data and local outputs remain outside Git",
    not bool(tracked_raw),
    tracked_raw if tracked_raw else "none tracked",
)


status = git(
    "status",
    "--short",
).stdout.strip()

if status:
    WARN(
        "Working tree is not clean",
        "This is acceptable for new paper figures/scripts, "
        "but commit them before submission.",
    )

    print(status)
else:
    PASS("Working tree clean")


# =====================================================================
# 2. Frozen detector provenance
# =====================================================================

print("\n--- 2. DETECTOR PROVENANCE ---")

weight = Path("yolov8n.pt")

if require_file(weight):
    check(
        "YOLOv8n weight hash",
        sha256(weight) == EXPECTED["yolo_weight_sha256"],
        sha256(weight),
    )


cubberly_pred = Path(
    "outputs/local/cubberly_native_full.jsonl"
)

if require_file(cubberly_pred):
    check(
        "Frozen Cubberly prediction hash",
        sha256(cubberly_pred)
        == EXPECTED["cubberly_prediction_sha256"],
        sha256(cubberly_pred),
    )


mem_root = Path(
    "outputs/local/memorial_court_holdout"
)

mem_pred = mem_root / "memorial_native_full.jsonl"
mem_det_summary_path = (
    mem_root / "memorial_native_detection_summary.json"
)

if require_file(mem_pred):
    check(
        "Memorial prediction hash",
        sha256(mem_pred)
        == EXPECTED["memorial_prediction_sha256"],
        sha256(mem_pred),
    )

if require_file(mem_det_summary_path):
    d = read_json(mem_det_summary_path)

    check(
        "Memorial detector sequence",
        d.get("sequence")
        == "memorial-court-2019-03-16_0",
    )

    check(
        "Memorial detector role",
        d.get("role")
        == "holdout_sensor_only_unscored",
    )

    check(
        "Memorial detector frame count",
        d.get("frames_count") == 1089,
        str(d.get("frames_count")),
    )

    check(
        "Memorial detector five-camera row count",
        d.get("processed_camera_frames") == 5445,
        str(d.get("processed_camera_frames")),
    )

    settings = d.get("settings", {})

    detector_expected = {
        "imgsz": 640,
        "confidence": 0.10,
        "iou": 0.70,
        "class_id": 0,
        "batch_size": 1,
        "stream_order": "frame then camera",
    }

    check(
        "Memorial detector settings frozen",
        all(
            settings.get(k) == v
            for k, v in detector_expected.items()
        ),
        str(settings),
    )

    check(
        "Memorial detector uses frozen YOLO weights",
        d.get("weight_sha256")
        == EXPECTED["yolo_weight_sha256"],
        str(d.get("weight_sha256")),
    )


# =====================================================================
# 3. Human blinded-review integrity
# =====================================================================

print("\n--- 3. HUMAN BLINDED REVIEW ---")

reviewer = Path(
    "research_records/"
    "cubberly_blinded_review/"
    "reviewer_a_completed.xlsx"
)

if require_file(reviewer):
    check(
        "Official completed reviewer workbook hash",
        sha256(reviewer)
        == EXPECTED["reviewer_xlsx_sha256"],
        sha256(reviewer),
    )

    try:
        from openpyxl import load_workbook

        wb = load_workbook(
            reviewer,
            read_only=True,
            data_only=True,
        )

        ws = (
            wb["in"]
            if "in" in wb.sheetnames
            else wb.active
        )

        rows = list(
            ws.iter_rows(values_only=True)
        )

        headers = [
            str(x).strip()
            if x is not None
            else ""
            for x in rows[0]
        ]

        records = [
            dict(zip(headers, row))
            for row in rows[1:]
            if any(v is not None for v in row)
        ]

        check(
            "Blinded review contains exactly 66 cases",
            len(records) == 66,
            str(len(records)),
        )

        ids = [
            str(r["case_id"]).strip()
            for r in records
        ]

        check(
            "Blinded review IDs are exactly D001-D066",
            ids
            == [
                f"D{i:03d}"
                for i in range(1, 67)
            ],
        )

        categorical = {
            "person_visible_in_box":
                {"yes", "no", "uncertain"},

            "association_clear":
                {"yes", "no", "uncertain"},

            "foreground_surface_clear":
                {"yes", "no", "uncertain"},

            "occlusion_or_edge_issue":
                {"yes", "no", "uncertain"},

            "reviewer_confidence":
                {"high", "medium", "low"},
        }

        normalized = {}

        all_valid = True

        for field, allowed in categorical.items():
            vals = [
                str(r[field]).strip().lower()
                if r[field] is not None
                else ""
                for r in records
            ]

            normalized[field] = vals

            if (
                any(not v for v in vals)
                or any(v not in allowed for v in vals)
            ):
                all_valid = False

        check(
            "All blinded-review judgments complete and valid",
            all_valid,
        )

        expected_counts = {
            "person_visible_in_box":
                Counter({
                    "yes": 59,
                    "uncertain": 7,
                }),

            "association_clear":
                Counter({
                    "yes": 37,
                    "uncertain": 29,
                }),

            "foreground_surface_clear":
                Counter({
                    "yes": 43,
                    "uncertain": 23,
                }),

            "occlusion_or_edge_issue":
                Counter({
                    "yes": 40,
                    "no": 26,
                }),

            "reviewer_confidence":
                Counter({
                    "high": 17,
                    "medium": 32,
                    "low": 17,
                }),
        }

        for field, expected_count in expected_counts.items():
            actual = Counter(
                normalized[field]
            )

            check(
                f"Blinded review count: {field}",
                actual == expected_count,
                str(dict(actual)),
            )

        uncertain_assoc_edge = sum(
            1
            for i in range(66)
            if (
                normalized["association_clear"][i]
                == "uncertain"
                and normalized["occlusion_or_edge_issue"][i]
                == "yes"
            )
        )

        check(
            "Uncertain associations concentrated in occlusion/edge cases",
            uncertain_assoc_edge == 28,
            f"{uncertain_assoc_edge}/29",
        )

    except ImportError:
        WARN(
            "openpyxl unavailable; workbook content counts not rechecked",
        )


# =====================================================================
# 4. Holdout sensor staging
# =====================================================================

print("\n--- 4. MEMORIAL COURT SENSOR STAGING ---")

sensor_manifest_path = Path(
    "manifests/holdout_memorial_sensor_only.json"
)

if require_file(sensor_manifest_path):
    m = read_json(sensor_manifest_path)

    check(
        "Holdout sequence identity",
        m.get("sequence")
        == "memorial-court-2019-03-16_0",
    )

    check(
        "Holdout sensor role",
        m.get("role")
        == "holdout_sensor_only_unscored",
    )

    check(
        "Sensor extraction completed",
        m.get("extracted") is True,
    )

    expected_streams = {
        "image_stitched",
        "image_0",
        "image_2",
        "image_4",
        "image_6",
        "image_8",
        "upper_lidar",
        "lower_lidar",
    }

    counts = m.get("frame_counts", {})

    check(
        "All expected Memorial streams present",
        set(counts) == expected_streams,
        str(sorted(counts)),
    )

    check(
        "Every Memorial sensor stream has 1089 frames",
        counts
        and all(v == 1089 for v in counts.values()),
        str(counts),
    )

    check(
        "Sensor frame IDs matched before scoring",
        m.get("all_sensor_frame_ids_match") is True,
    )

    check(
        "Reference archive unopened during sensor staging",
        m.get("reference_archive_opened") is False
        and m.get("reference_annotations_inspected") is False,
    )


# =====================================================================
# 5. Frozen range-method equivalence
# =====================================================================

print("\n--- 5. RANGE METHOD FREEZE ---")

dev_sensor = Path(
    "scripts/cache_jrdb_dev_sensor_depth_candidates.py"
)

holdout_sensor = Path(
    "scripts/cache_jrdb_holdout_sensor_depth_candidates.py"
)

expected_constants = {
    "BIN_WIDTH_M": 0.5,
    "MAX_RANGE_M": 20.0,
    "MIN_BIN_POINTS": 3,
    "MIN_BAND_POINTS": 8,
    "DOMINANCE_RATIO": 2.0,
}

if (
    require_file(dev_sensor)
    and require_file(holdout_sensor)
):
    for name, value in expected_constants.items():
        a = literal_constant(
            dev_sensor,
            name,
        )

        b = literal_constant(
            holdout_sensor,
            name,
        )

        check(
            f"Frozen range constant {name}",
            a == value and b == value,
            f"dev={a}, holdout={b}",
        )

    for fn in (
        "project_ego_xy",
        "surface_bands",
        "evaluate_box",
    ):
        check(
            f"Holdout uses identical frozen function: {fn}",
            function_ast(dev_sensor, fn)
            == function_ast(holdout_sensor, fn),
        )


# =====================================================================
# 6. Sensor candidate output
# =====================================================================

print("\n--- 6. MEMORIAL SENSOR CANDIDATES ---")

candidate_csv = (
    mem_root / "memorial_sensor_surface_full.csv"
)

candidate_summary_path = (
    mem_root / "memorial_sensor_surface_summary.json"
)

if require_file(candidate_csv):
    check(
        "Memorial candidate CSV hash",
        sha256(candidate_csv)
        == EXPECTED["memorial_candidate_sha256"],
        sha256(candidate_csv),
    )

if require_file(candidate_summary_path):
    s = read_json(candidate_summary_path)

    check(
        "Candidate sequence",
        s.get("sequence")
        == "memorial-court-2019-03-16_0",
    )

    check(
        "Candidate output still marked unvalidated sensor-side",
        s.get("status")
        == "unvalidated_measured_surface_candidates",
    )

    check(
        "Candidate method role remained sensor-only",
        s.get("role")
        == "holdout_sensor_only_unscored",
    )

    check(
        "Candidate detector hash linkage",
        s.get("raw_prediction_sha256")
        == EXPECTED["memorial_prediction_sha256"],
    )

    settings = s.get("settings", {})

    check(
        "Candidate numerical settings match freeze",
        settings == {
            "bin_width_m": 0.5,
            "min_points_per_bin": 3,
            "min_points_per_candidate_band": 8,
            "dominance_ratio": 2.0,
            "max_ego_xy_radius_m": 20.0,
        },
        str(settings),
    )

    status_counts = s.get(
        "status_counts",
        {}
    )

    expected_status = {
        "single_supported_surface_candidate": 11502,
        "dominant_surface_candidate": 6098,
        "insufficient_supported_surface": 6554,
        "ambiguous_surface_abstain": 838,
    }

    check(
        "Sensor candidate status counts",
        status_counts == expected_status,
        str(status_counts),
    )

    check(
        "Sensor candidate totals sum to 24,992 detector boxes",
        sum(status_counts.values()) == 24992
        and s.get("boxes") == 24992,
        str(sum(status_counts.values())),
    )


# =====================================================================
# 7. Pre-reference output freeze
# =====================================================================

print("\n--- 7. PRE-REFERENCE OUTPUT FREEZE ---")

pre_ref_path = Path(
    "manifests/"
    "memorial_court_sensor_output_freeze.json"
)

if require_file(pre_ref_path):
    p = read_json(pre_ref_path)

    check(
        "Sensor outputs were frozen before reference opening",
        p.get("role")
        == "sensor_outputs_frozen_before_reference_opening"
        and p.get("reference_archive_opened") is False
        and p.get("reference_annotations_inspected") is False,
    )

    frozen_files = p.get("files", {})

    check(
        "Frozen detector hash recorded correctly",
        frozen_files.get(
            "detections", {}
        ).get("sha256")
        == EXPECTED["memorial_prediction_sha256"],
    )

    check(
        "Frozen candidate hash recorded correctly",
        frozen_files.get(
            "surface_candidates", {}
        ).get("sha256")
        == EXPECTED["memorial_candidate_sha256"],
    )


# =====================================================================
# 8. Frozen evaluator implementation
# =====================================================================

print("\n--- 8. HOLDOUT EVALUATOR FREEZE ---")

dev_eval = Path(
    "scripts/evaluate_jrdb_dev_surface_candidates_offline.py"
)

hold_eval = Path(
    "scripts/evaluate_jrdb_holdout_surface_candidates_offline.py"
)

if require_file(dev_eval) and require_file(hold_eval):
    for name, expected in {
        "CAMERAS": (0, 2, 4, 6, 8),
        "MIN_IOU": 0.5,
        "NEAR_RADIUS_M": 3.0,
    }.items():

        a = literal_constant(
            dev_eval,
            name,
        )

        b = literal_constant(
            hold_eval,
            name,
        )

        check(
            f"Frozen evaluator constant {name}",
            a == expected and b == expected,
            f"dev={a}, holdout={b}",
        )

    for fn in (
        "optional_float",
        "match_one_to_one",
        "rate_stats",
    ):
        check(
            f"Holdout evaluator uses frozen function: {fn}",
            function_ast(dev_eval, fn)
            == function_ast(hold_eval, fn),
        )


scoring_manifest_path = Path(
    "manifests/holdout_memorial_scoring.json"
)

if require_file(scoring_manifest_path):
    sm = read_json(scoring_manifest_path)

    check(
        "Scoring manifest tied to method freeze",
        sm.get("method_freeze_commit")
        == FREEZE_COMMIT,
    )

    check(
        "Scoring IoU frozen at 0.5",
        float(
            sm.get("matching_2d_iou_threshold")
        ) == 0.5,
    )

    check(
        "Scoring range boundary frozen at 3.0 m",
        float(sm.get("near_radius_m"))
        == 3.0,
    )

    check(
        "Post-reference tuning explicitly prohibited",
        sm.get("post_reference_tuning_allowed")
        is False,
    )


# =====================================================================
# 9. Holdout scoring output
# =====================================================================

print("\n--- 9. MEMORIAL HOLDOUT SCORING ---")

eval_summary_path = (
    mem_root / "memorial_holdout_evaluation_summary.json"
)

eval_csv = (
    mem_root / "memorial_surface_reference_holdout.csv"
)

if require_file(eval_summary_path):
    e = read_json(eval_summary_path)

    check(
        "Evaluation status is frozen holdout",
        e.get("status")
        == "frozen_holdout_annotation_center_comparison",
    )

    check(
        "Evaluation role is frozen holdout scoring",
        e.get("role")
        == "frozen_holdout_scoring",
    )

    check(
        "Frozen candidate CSV used for scoring",
        e.get("candidate_csv_sha256")
        == EXPECTED["memorial_candidate_sha256"],
    )

    if require_file(candidate_summary_path):
        check(
            "Candidate summary hash linked to evaluation",
            e.get("candidate_summary_sha256")
            == sha256(candidate_summary_path),
            str(e.get("candidate_summary_sha256")),
        )

    check(
        "Memorial 3D reference hash",
        e.get("label_3d_sha256")
        == EXPECTED["memorial_label_3d_sha256"],
        str(e.get("label_3d_sha256")),
    )

    check(
        "Five native 2D reference members hash",
        e.get("five_2d_members_combined_sha256")
        == EXPECTED["memorial_2d_combined_sha256"],
        str(
            e.get(
                "five_2d_members_combined_sha256"
            )
        ),
    )

    counts = e.get("counts", {})

    expected_counts = {
        "detector_boxes": 24992,
        "evaluable_2d_labels": 18390,
        "matched_2d_boxes": 13349,
        "matched_2d_and_evaluable_3d": 13246,
        "2d_no_eval_or_unknown": 7476,
        "matches_without_evaluable_3d": 103,
    }

    check(
        "Evaluation population counts",
        all(
            counts.get(k) == v
            for k, v in expected_counts.items()
        ),
        str(counts),
    )

    check(
        "2D matched minus missing-3D equals scored population",
        counts.get("matched_2d_boxes", -1)
        - counts.get(
            "matches_without_evaluable_3d",
            -999,
        )
        == counts.get(
            "matched_2d_and_evaluable_3d",
            -2,
        ),
    )

    all3d = e.get(
        "all_matched_evaluable_3d",
        {},
    )

    check(
        "Matched evaluable 3D population",
        all3d.get(
            "matched_reference_pairs"
        ) == 13246,
    )

    check(
        "Candidate available pairs",
        all3d.get(
            "candidate_surface", {}
        ).get("available_pairs")
        == 12354,
    )

    check(
        "Nearest available pairs",
        all3d.get(
            "nearest_supported_surface", {}
        ).get("available_pairs")
        == 12992,
    )

    check(
        "Holdout median discrepancies",
        all3d.get(
            "candidate_surface", {}
        ).get(
            "median_absolute_center_discrepancy_m"
        ) == 0.076
        and all3d.get(
            "nearest_supported_surface", {}
        ).get(
            "median_absolute_center_discrepancy_m"
        ) == 0.09,
    )

    check(
        "Holdout mean discrepancies",
        all3d.get(
            "candidate_surface", {}
        ).get(
            "mean_absolute_center_discrepancy_m"
        ) == 0.176
        and all3d.get(
            "nearest_supported_surface", {}
        ).get(
            "mean_absolute_center_discrepancy_m"
        ) == 0.931,
    )

    check(
        ">1 m discrepancy counts",
        all3d.get(
            "candidate_surface", {}
        ).get(
            "absolute_discrepancies_over_1m"
        ) == 313
        and all3d.get(
            "nearest_supported_surface", {}
        ).get(
            "absolute_discrepancies_over_1m"
        ) == 2384,
    )


# =====================================================================
# 10. Derived final reporting
# =====================================================================

print("\n--- 10. FINAL REPORTING ARITHMETIC ---")

report_path = (
    mem_root / "memorial_holdout_reporting.json"
)

if require_file(report_path):
    r = read_json(report_path)

    check(
        "Final reporting JSON hash",
        sha256(report_path)
        == EXPECTED["memorial_reporting_sha256"],
        sha256(report_path),
    )

    if require_file(eval_csv):
        check(
            "Reporting file linked to exact scored CSV",
            r.get("source_csv_sha256")
            == sha256(eval_csv),
            sha256(eval_csv),
        )

    check(
        "Final scored population",
        r.get("matched_evaluable_3d_pairs")
        == 13246,
    )

    ref = r.get(
        "reference_distribution",
        {},
    )

    check(
        "Reference near/far counts",
        ref.get("near_or_equal_3m") == 842
        and ref.get("farther_than_3m") == 12404
        and 842 + 12404 == 13246,
        str(ref),
    )

    cov = r.get("coverage", {})

    check(
        "Candidate coverage counts",
        cov.get("candidate_available")
        == 12354
        and cov.get(
            "candidate_abstain_or_unavailable"
        ) == 892
        and 12354 + 892 == 13246,
    )

    check(
        "Nearest coverage counts",
        cov.get("nearest_available")
        == 12992
        and cov.get("nearest_unavailable")
        == 254
        and 12992 + 254 == 13246,
    )

    c = r.get(
        "candidate_3m_decisions",
        {},
    )

    n = r.get(
        "nearest_3m_decisions",
        {},
    )

    check(
        "Depth-band 3 m confusion counts",
        c.get(
            "reference_near__predicted_near"
        ) == 820
        and c.get(
            "reference_near__predicted_far"
        ) == 10
        and c.get(
            "reference_far__predicted_near"
        ) == 179
        and c.get(
            "reference_far__predicted_far"
        ) == 11345
        and c.get(
            "abstain_or_unavailable"
        ) == 892,
    )

    check(
        "Depth-band available-decision arithmetic",
        820 + 10 + 179 + 11345
        == c.get("available")
        == 12354,
    )

    check(
        "Nearest 3 m confusion counts",
        n.get(
            "reference_near__predicted_near"
        ) == 835
        and n.get(
            "reference_near__predicted_far"
        ) == 7
        and n.get(
            "reference_far__predicted_near"
        ) == 422
        and n.get(
            "reference_far__predicted_far"
        ) == 11728
        and n.get(
            "abstain_or_unavailable"
        ) == 254,
    )

    check(
        "Nearest available-decision arithmetic",
        835 + 7 + 422 + 11728
        == n.get("available")
        == 12992,
    )

    paired = r.get(
        "paired_range_discrepancy",
        {},
    )

    check(
        "Paired discrepancy population",
        paired.get("pairs") == 12354,
    )

    check(
        "Paired win/tie arithmetic",
        paired.get(
            "candidate_lower_discrepancy"
        ) == 1841
        and paired.get(
            "nearest_lower_discrepancy"
        ) == 5
        and paired.get(
            "equal_discrepancy"
        ) == 10508
        and 1841 + 5 + 10508 == 12354,
    )

    check(
        "Paired discrepancy values",
        paired.get(
            "candidate_median_abs_center_discrepancy_m"
        ) == 0.0763
        and paired.get(
            "nearest_median_abs_center_discrepancy_m"
        ) == 0.0874
        and paired.get(
            "candidate_mean_abs_center_discrepancy_m"
        ) == 0.1757
        and paired.get(
            "nearest_mean_abs_center_discrepancy_m"
        ) == 0.8527,
    )

    pd = r.get(
        "paired_3m_decision_comparison",
        {},
    )

    check(
        "Paired 3 m agreement arithmetic",
        pd.get(
            "both_agree_with_reference"
        ) == 11940
        and pd.get(
            "candidate_only_agrees"
        ) == 225
        and pd.get(
            "nearest_only_agrees"
        ) == 3
        and pd.get(
            "neither_agrees"
        ) == 186
        and 11940 + 225 + 3 + 186
        == 12354,
    )

    check(
        "All 228 method disagreements are candidate-far/nearest-near",
        pd.get(
            "candidate_far__nearest_near"
        ) == 228
        and pd.get(
            "candidate_only_agrees"
        )
        + pd.get(
            "nearest_only_agrees"
        )
        == 228,
    )

    linkage = r.get(
        "detector_linkage",
        {},
    )

    check(
        "Detector linkage arithmetic",
        linkage.get("detector_boxes")
        == 24992
        and linkage.get("matched_2d_boxes")
        == 13349
        and linkage.get(
            "unlinked_detector_boxes"
        ) == 11643
        and 13349 + 11643 == 24992,
    )

    check(
        "Reference linkage arithmetic",
        linkage.get(
            "evaluable_2d_labels"
        ) == 18390
        and linkage.get(
            "unmatched_evaluable_2d_labels"
        ) == 5041
        and 13349 + 5041 == 18390,
    )


# =====================================================================
# 11. Paper table artifacts
# =====================================================================

print("\n--- 11. PAPER TABLE ARTIFACTS ---")

table_paths = [
    Path(
        "paper/generated/"
        "memorial_table1_range_comparison.csv"
    ),
    Path(
        "paper/generated/"
        "memorial_table2_3m_decisions.csv"
    ),
    Path(
        "paper/generated/"
        "memorial_table3_paired_comparison.csv"
    ),
    Path(
        "paper/generated/"
        "memorial_table4_linkage.csv"
    ),
    Path(
        "paper/generated/"
        "memorial_holdout_tables.md"
    ),
]

for p in table_paths:
    if p.exists():
        PASS(
            "Paper table artifact exists",
            str(p),
        )
    else:
        WARN(
            "Paper table artifact missing",
            str(p),
        )


# =====================================================================
# 12. Result-strength summary
# =====================================================================

print("\n--- 12. SCIENTIFIC RESULT SUMMARY ---")

candidate_coverage = (
    12354 / 13246 * 100
)

nearest_coverage = (
    12992 / 13246 * 100
)

false_near_reduction = (
    (422 - 179) / 422 * 100
)

print(
    f"Depth-band coverage: "
    f"{candidate_coverage:.2f}%"
)

print(
    f"Nearest coverage: "
    f"{nearest_coverage:.2f}%"
)

print(
    "Far-reference -> near decisions: "
    "179 depth-band vs 422 nearest"
)

print(
    f"Raw reduction in far->near counts: "
    f"{false_near_reduction:.1f}% "
    "(interpret with coverage difference)"
)

print(
    "Paired discrepancy comparison: "
    "depth-band better 1841, nearest better 5, "
    "equal 10508"
)

print(
    "Paired 3 m disagreements: "
    "225 favor depth-band, 3 favor nearest"
)

print(
    "Near-reference decisions: "
    "depth-band 820 near / 10 far / 12 unavailable; "
    "nearest 835 near / 7 far / 0 unavailable"
)

print(
    "Human review: 66 enriched diagnostic cases; "
    "descriptive only, not an accuracy estimate."
)


# =====================================================================
# 13. Manual claim-safety checklist
# =====================================================================

print("\n--- 13. CLAIM-SAFETY CHECKLIST ---")

claim_rules = [
    "Do NOT call 0.076 m a direct person-distance error; it is "
    "surface-to-annotated-center discrepancy.",

    "Do NOT report 225/228 as general accuracy; it applies only "
    "to paired cases where the methods disagreed.",

    "Do NOT call unlinked detector boxes false positives without "
    "independent annotation establishing that.",

    "Do NOT convert the enriched 66-case blinded review into an "
    "accuracy percentage.",

    "Do NOT claim temporal alert accuracy, event latency, tracker "
    "performance, or robot end-to-end latency.",

    "State that adjacent frames/cameras/repeated persons are "
    "statistically dependent.",

    "State that only one location-separated independent holdout "
    "recording was scored.",

    "State the coverage trade-off: depth-band abstains more often.",

    "Mention the strong class imbalance: 842 near references vs "
    "12,404 far references.",

    "Mention camera-specific limitations, including no matched "
    "<=3 m references for camera 2 in the holdout.",
]

for rule in claim_rules:
    print("[MANUAL]", rule)


# =====================================================================
# FINAL VERDICT
# =====================================================================

print("\n" + "=" * 78)
print("FINAL VERDICT")
print("=" * 78)

print(
    f"PASS checks: {len(passes)}"
)

print(
    f"WARNINGS:    {len(warnings)}"
)

print(
    f"FAILURES:    {len(failures)}"
)

if failures:
    print()
    print(
        "EXPERIMENT CLOSURE: FAIL"
    )

    print(
        "Do not finalize the manuscript until the failures above "
        "are understood. Do not rerun/tune automatically."
    )

    sys.exit(1)

print()
print(
    "EXPERIMENT CLOSURE: PASS"
)

print(
    "The frozen development/holdout pipeline is internally "
    "consistent and the completed experimental results can now "
    "be treated as final."
)

print()
print(
    "RESULT INTERPRETATION: SUPPORTED"
)

print(
    "The evidence supports the narrow claim that supported "
    "depth-band selection suppresses many nuisance-induced "
    "near-range assignments relative to the nearest-supported-"
    "surface baseline, while trading some range coverage for "
    "conservative abstention."
)

print()
print(
    "This audit does NOT establish publication acceptance, "
    "independent-observation confidence intervals, robot latency, "
    "or temporal-event performance."
)

report_dir = Path("paper/generated")
report_dir.mkdir(
    parents=True,
    exist_ok=True,
)

report = {
    "experiment_closure":
        "PASS",

    "passes":
        len(passes),

    "warnings":
        len(warnings),

    "failures":
        len(failures),

    "method_freeze_tag":
        FREEZE_TAG,

    "method_freeze_commit":
        FREEZE_COMMIT,

    "memorial_prediction_sha256":
        EXPECTED[
            "memorial_prediction_sha256"
        ],

    "memorial_candidate_sha256":
        EXPECTED[
            "memorial_candidate_sha256"
        ],

    "memorial_reporting_sha256":
        EXPECTED[
            "memorial_reporting_sha256"
        ],

    "final_narrow_claim":
        "Supported depth-band selection reduces nuisance-induced "
        "near-range assignments relative to nearest-supported-"
        "surface selection, at the cost of lower coverage and "
        "more abstention.",

    "limitations": claim_rules,
}

target = (
    report_dir
    / "final_experiment_sanity.json"
)

target.write_text(
    json.dumps(
        report,
        indent=2,
    ) + "\n",
    encoding="utf-8",
)

print()
print(
    "Saved machine-readable closure report:",
    target,
)
